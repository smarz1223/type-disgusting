"""
Type Disgusting Fantasy Hockey - data pipeline
Runs daily on GitHub Actions. Downloads the two published Google Sheets
workbooks, scores every player from the Yahoo team logs, pairs weekly
matchups from Free For All, builds standings, analytics, and league history.

Outputs:
  data/td_data.json       everything the site needs
  data/recon_report.md    human-readable reconciliation report

Local test: set env vars STATS_FILE and HISTORY_FILE to .xlsx paths.
"""
import io, json, math, os, re, collections, datetime
import openpyxl

# ----------------------------------------------------------------- CONFIG
STATS_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vRJ7-9hkb-6d3uoOfyCu0T9xrwpWweOGXoL9FgXaqEHv40SeH2-Z4YpRbT_pA9-06s0lmUSmYniNvdc/pub?output=xlsx")
HISTORY_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vSA7UMmJkxU9_22-hqHj51aPMDMyxxwCXqPV1eTEThSkfSIcWe-2sXq8m5CssbiY7lGW5mnKZdfdsT9/pub?output=xlsx")

SEASON = 2026                  # season start year; displayed as 2026-27
REG_SEASON_WEEKS = 23
PLAYOFF_WEEKS = [24, 25]
PLAYOFF_TEAMS = 4
OUT_DIR = "data"
YAHOO_LEAGUE_URL = "https://hockey.fantasysports.yahoo.com/hockey/24054"

# 2026-27 owners (display names, used everywhere on the site)
OWNERS = ["MARZ", "PHIL", "ADAM B", "PEDRO", "PAT CLARK", "SCOTTY", "FUR", "MEANY"]
# Yahoo team names -> owners. Free For All and Data Pull use team names.
# If a team renames itself, add the new name here (old names can stay).
TEAM_TO_OWNER = {
    "Cat Plark": "PHIL",
    "Soft Dump in Corner": "MARZ",
    "Wild Hogs": "ADAM B",
    "Panarin Bread": "PEDRO",
    "H DUBZ ON THA TRACK": "PAT CLARK",
    "Bros Before Hossas": "SCOTTY",
    "Ladies...Beverages": "FUR",
    "Sloppy Seconds": "MEANY",
}
# History names -> display names (everyone else: current owners uppercased, former as-is)
HISTORY_NAME_MAP = {"Scott Ratto": "SCOTTY"}

# League scoring (Yahoo league settings, captured Oct 8, 2026)
SKATER_SCORING = {"G": 15, "A": 10, "+/-": 5, "PIM": 2.5, "PPP": 5, "SOG": 2}
GOALIE_SCORING = {"W": 20, "GA": -5, "SV": 1, "SHO": 10}
SCORING = {**SKATER_SCORING, **GOALIE_SCORING}
CAT_LABELS = {"G": "Goals", "A": "Assists", "+/-": "Plus/Minus", "PIM": "Penalty Minutes",
              "PPP": "Powerplay Points", "SOG": "Shots on Goal", "W": "Wins",
              "GA": "Goals Against", "SV": "Saves", "SHO": "Shutouts"}
CATEGORY_GROUPS = {"Skaters": list(SKATER_SCORING), "Goalies": list(GOALIE_SCORING)}
POSITIONS = ["F", "D", "G"]

# Yahoo name-cell cleanup
NOTE_FLAGS = ["No new player Notes", "No new player Note", "New Player Notes", "New Player Note",
              "Player Notes", "Player Note", "Video Forecast"]
STATUS_TAGS = sorted(["IR-LT", "IR-NR", "IR", "DTD", "SUSP", "INJ", "NA", "O"], key=len, reverse=True)
TEAM_POS_RE = re.compile(r"([A-Za-z]{2,3}) - ((?:C|LW|RW|D|G|Util)(?:,(?:C|LW|RW|D|G|Util))*)\s*$")


# ----------------------------------------------------------------- HELPERS
def load_workbook(env_var, url):
    path = os.environ.get(env_var)
    if path:
        return openpyxl.load_workbook(path, data_only=True)
    import requests
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def r2(x):
    return round(x + 0.0, 2)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


_TEAM_IDX = {norm(k): v for k, v in TEAM_TO_OWNER.items()}
_TEAM_IDX.update({norm(o): o for o in OWNERS})


def to_owner(name):
    return _TEAM_IDX.get(norm(name))


def season_label(y):
    return f"{y}-{str(y + 1)[-2:]}"


def parse_name(raw):
    first = str(raw).split("\n")[0].strip()
    m = TEAM_POS_RE.search(first)
    team, elig = (m.group(1).upper(), m.group(2)) if m else (None, "")
    name = first[:m.start()] if m else first
    for f in NOTE_FLAGS:
        name = name.replace(f, "")
    name = name.strip()
    for tag in STATUS_TAGS:  # status glued to the name, e.g. "Aleksander BarkovIR-LT"
        if name.endswith(tag) and len(name) > len(tag):
            prev = name[-len(tag) - 1]
            if prev.islower() or prev in ".'":
                name = name[:-len(tag)].strip()
                break
    return name, team, elig


def pos_group(elig, kind):
    """Forward (C/LW/RW), Defenseman, or Goalie."""
    if kind == "G":
        return "G"
    parts = [p for p in elig.split(",") if p in ("C", "LW", "RW", "D")]
    if not parts:
        return None
    return "D" if parts[0] == "D" else "F"


# ----------------------------------------------------------------- TEAM LOGS
def read_pull(ws, scoring, kind, flags):
    """SKATER PULL / GOALIE PULL: one 12-column block per owner, side by side.
    Returns (player_rows, yahoo_totals_by_owner)."""
    rows = list(ws.iter_rows(values_only=True))
    top = rows[0]
    starts = []
    for i, v in enumerate(top):
        if v and to_owner(v) and (i == 0 or top[i - 1] != v):
            starts.append((i, to_owner(v)))
    if not starts:
        raise RuntimeError(f"{ws.title}: no owner blocks found in row 1")
    players, totals = [], {}
    for n, (s, owner) in enumerate(starts):
        e = starts[n + 1][0] if n + 1 < len(starts) else len(top)
        hdr = [str(v).strip() if v is not None else None for v in rows[1][s:e]]
        idx = {}
        for j, h in enumerate(hdr):
            if h == "Name":
                idx["Name"] = s + j
            elif h == "GP*":
                idx["GP"] = s + j
            elif h in scoring and h not in idx:
                idx[h] = s + j
        missing = {"Name", "GP", *scoring} - set(idx)
        if missing:
            flags.append(f"{ws.title} {owner}: columns not found {sorted(missing)}")
            continue
        for r in rows[2:]:
            name = r[idx["Name"]] if idx["Name"] < len(r) else None
            if name in (None, ""):
                continue
            stats = {k: num(r[i]) for k, i in idx.items() if k != "Name"}
            if str(name).strip() == "Totals":
                totals[owner] = stats
                continue
            pname, team, elig = parse_name(name)
            players.append({"owner": owner, "player": pname, "team": team, "elig": elig,
                            "pos": pos_group(elig, kind), "kind": kind, "stats": stats})
    return players, totals


def score(stats, scoring):
    return {k: stats.get(k, 0) * v for k, v in scoring.items()}


# ----------------------------------------------------------------- WEEKLY
def load_weekly(ws, flags):
    weeks = collections.defaultdict(dict)
    unknown = set()
    for r in ws.iter_rows(min_row=3, values_only=True):
        if len(r) < 4:
            continue
        wk, team, pf, pa = r[0], r[1], r[2], r[3]
        if wk in (None, "") or not team or pf in (None, "") or pa in (None, ""):
            continue
        owner = to_owner(team)
        if not owner:
            unknown.add(str(team).strip())
            continue
        weeks[int(num(wk))][owner] = {"pf": num(pf), "pa": num(pa)}
    for u in sorted(unknown):
        flags.append(f"Free For All team '{u}' is not in TEAM_TO_OWNER. Add it to fetch_data.py.")
    return dict(sorted(weeks.items()))


def pair_opponents(weeks, flags):
    for wk, teams in weeks.items():
        for t, d in teams.items():
            cands = [o for o, e in teams.items() if o != t
                     and abs(e["pf"] - d["pa"]) < 0.005 and abs(e["pa"] - d["pf"]) < 0.005]
            if len(cands) == 1:
                d["opp"] = cands[0]
            else:
                d["opp"] = None
                flags.append(f"Week {wk}: could not uniquely match opponent for {t} ({len(cands)} candidates)")


def result(a, b):
    return "W" if a > b + 1e-9 else ("L" if a < b - 1e-9 else "T")


def norm_cdf(z):
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def build_standings(weeks, owners):
    rec = {o: collections.Counter() for o in owners}
    weekly_rows, perf = [], collections.defaultdict(list)
    for wk, teams in weeks.items():
        scores = [d["pf"] for d in teams.values()]
        n = len(scores)
        mean = sum(scores) / n
        sd = math.sqrt(sum((s - mean) ** 2 for s in scores) / n) or 1
        reg = wk <= REG_SEASON_WEEKS
        for t, d in teams.items():
            h2h = result(d["pf"], d["pa"])
            rank = 1 + sum(1 for s in scores if s > d["pf"] + 1e-9)
            ap = collections.Counter(result(d["pf"], e["pf"]) for o, e in teams.items() if o != t)
            weekly_rows.append({"week": wk, "team": t, "opp": d.get("opp"), "pf": r2(d["pf"]),
                                "pa": r2(d["pa"]), "margin": r2(d["pf"] - d["pa"]), "h2h": h2h,
                                "rank": rank, "allplay": f"{ap['W']}-{ap['L']}-{ap['T']}",
                                "avg": r2(mean), "playoff": not reg})
            if not reg:
                continue
            c = rec[t]
            c[h2h] += 1
            c["AP_W"] += ap["W"]; c["AP_L"] += ap["L"]; c["AP_T"] += ap["T"]
            c["PF"] += d["pf"]; c["PA"] += d["pa"]; c["G"] += 1
            if rank == 1:
                c["HIGH"] += 1
            perf[t].append(norm_cdf((d["pf"] - mean) / sd) * 100)
    table = []
    for t in owners:
        c = rec[t]
        g = c["W"] + c["L"] + c["T"]
        ap_g = c["AP_W"] + c["AP_L"] + c["AP_T"]
        pct = (c["W"] + 0.5 * c["T"]) / g if g else 0
        ap_pct = (c["AP_W"] + 0.5 * c["AP_T"]) / ap_g if ap_g else 0
        table.append({
            "team": t, "W": c["W"], "L": c["L"], "T": c["T"], "win_pct": round(pct, 4),
            "allplay": f"{c['AP_W']}-{c['AP_L']}-{c['AP_T']}", "allplay_pct": round(ap_pct, 4),
            "exp_wins": round(ap_pct * g, 2),                 # all-play expected wins
            "luck": round((pct - ap_pct + 1) * 50, 1),        # 0-100, 50 = neutral
            "perf_rating": round(sum(perf[t]) / len(perf[t]), 1) if perf[t] else None,
            "PF": r2(c["PF"]), "PA": r2(c["PA"]), "PPG": r2(c["PF"] / c["G"]) if c["G"] else 0,
            "G": c["G"], "weekly_highs": c["HIGH"],
        })
    table.sort(key=lambda x: (-x["win_pct"], -x["PF"]))
    for i, row in enumerate(table, 1):
        row["seed"] = i
        row["status"] = "playoff" if i <= PLAYOFF_TEAMS else "out"
    return table, weekly_rows


# ----------------------------------------------------------------- RECONCILIATION
def read_data_pull(ws):
    """Data Pull tab: Yahoo league stat totals through completed matchups."""
    blocks, cur = [], None
    for r in ws.iter_rows(values_only=True):
        if r[0] == "Team Name":
            cur = {"hdr": [str(v).strip() if v else None for v in r[:7]], "rows": {}}
            blocks.append(cur)
            continue
        if cur is None:
            continue
        if r[0] in (None, ""):
            cur = None
            continue
        cur["rows"][r[0]] = {h: num(v) for h, v in zip(cur["hdr"][1:], r[1:7]) if h}
    sk = next((b for b in blocks if "G" in b["hdr"]), None)
    gl = next((b for b in blocks if "W" in b["hdr"]), None)
    out = {}
    if not sk or not gl:
        return out
    for team in set(sk["rows"]) | set(gl["rows"]):
        o = to_owner(team)
        if not o:
            continue
        s, g = sk["rows"].get(team, {}), gl["rows"].get(team, {})
        out[o] = sum(s.get(k, 0) * v for k, v in SKATER_SCORING.items()) + \
            sum(g.get(k, 0) * v for k, v in GOALIE_SCORING.items())
    return out


def reconcile(players, yahoo_totals, weeks, owners, official, flags):
    # 1. parsed player rows vs Yahoo "Totals" rows
    for kind, tot in yahoo_totals.items():
        for o, ytot in tot.items():
            mine = collections.Counter()
            for p in players:
                if p["owner"] == o and p["kind"] == kind:
                    mine.update(p["stats"])
            for k, v in ytot.items():
                if abs(mine[k] - v) > 0.001:
                    flags.append(f"{o} {kind} log: parsed {k}={mine[k]} vs Yahoo Totals row {v}")
    weeks_entered = len(weeks)
    # 2. Yahoo league totals (Data Pull) vs cumulative Free For All scores
    verified = 0
    if official and weeks:
        wk_list = list(weeks)
        for k in range(len(wk_list), 0, -1):
            cum = {o: sum(weeks[w][o]["pf"] for w in wk_list[:k] if o in weeks[w]) for o in owners}
            if all(abs(cum[o] - official.get(o, -1e9)) < 0.011 for o in owners):
                verified = wk_list[k - 1]
                break
    # 3. team log (daily) vs entered scores
    report = []
    for o in owners:
        wk_total = sum(w[o]["pf"] for w in weeks.values() if o in w)
        calc = sum(p["points"] for p in players if p["owner"] == o)
        diff = calc - wk_total
        state = "MATCH" if abs(diff) < 0.011 else ("AHEAD" if diff > 0 else "BEHIND")
        report.append({"team": o, "weekly_total": r2(wk_total), "game_log": r2(calc),
                       "in_progress": r2(diff), "yahoo_official": r2(official[o]) if o in official else None,
                       "state": state})
    behind = [r["team"] for r in report if r["state"] == "BEHIND"]
    if behind:
        flags.append(f"Game log is below entered scores for {', '.join(behind)}. "
                     "IMPORTHTML may not have refreshed, or a score was typed wrong.")
    hard = any("Totals row" in f or "opponent" in f or "TEAM_TO_OWNER" in f for f in flags)
    if not weeks:
        status = "PRESEASON"
    elif hard:
        status = "FAIL"
    elif verified == list(weeks)[-1]:
        status = "PASS"
    elif verified:
        status = "PARTIAL"
        flags.append(f"Yahoo's league totals tie to Free For All through week {verified}; "
                     f"later weeks are entered but Yahoo's table hasn't caught up yet.")
    else:
        status = "CHECK"
        flags.append("Free For All scores don't tie to Yahoo's league stat totals (Data Pull tab) "
                     "for any week. Check for a typo in a weekly score.")
    return report, status, weeks_entered, verified


# ----------------------------------------------------------------- ANALYTICS
def build_analytics(players, owners):
    raw = {o: collections.Counter() for o in owners}
    pts = {o: collections.Counter() for o in owners}
    by_pos = {o: collections.Counter() for o in owners}
    for p in players:
        o = p["owner"]
        raw[o].update({k: v for k, v in p["stats"].items() if k != "GP"})
        pts[o].update(p["pts_by_cat"])
        if p["pos"]:
            by_pos[o][p["pos"]] += p["points"]
    cats = {o: {"raw": {k: raw[o].get(k, 0) for k in SCORING},
                "pts": {k: r2(pts[o].get(k, 0)) for k in SCORING},
                "groups": {g: r2(sum(pts[o].get(k, 0) for k in ks)) for g, ks in CATEGORY_GROUPS.items()},
                "total": r2(sum(pts[o].values()))} for o in owners}
    positions = {o: {pos: r2(by_pos[o].get(pos, 0)) for pos in POSITIONS} for o in owners}
    return cats, positions


# ----------------------------------------------------------------- HISTORY
def hname(n):
    n = str(n).strip()
    n = HISTORY_NAME_MAP.get(n, n)
    return n.upper() if n.upper() in OWNERS else n


def build_history(wb, flags):
    ws = wb["Data"]
    hdr = [str(v).strip() if v else None for v in next(ws.iter_rows(max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(hdr) if h}
    need = {"Years", "Manager", "1st", "2nd", "3rd", "Wins", "Losses", "Ties", "Games", "Playoffs"}
    if need - set(col):
        flags.append(f"History Data tab missing columns {sorted(need - set(col))}")
        return None
    seasons = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Years"]] in (None, "") or not r[col["Manager"]]:
            continue
        g = lambda k: num(r[col[k]])
        seasons.append({"year": int(num(r[col["Years"]])), "manager": hname(r[col["Manager"]]),
                        "first": g("1st"), "second": g("2nd"), "third": g("3rd"),
                        "W": g("Wins"), "L": g("Losses"), "T": g("Ties"),
                        "games": g("Games"), "playoffs": g("Playoffs")})
    # active = in this season's league (overrides the sheet's Status column)
    agg = collections.defaultdict(collections.Counter)
    for s in seasons:
        a = agg[s["manager"]]
        a["years"] += 1
        for k in ("first", "second", "third", "W", "L", "T", "games", "playoffs"):
            a[k] += s[k]
    alltime = []
    for m, a in agg.items():
        alltime.append({"manager": m, "active": m in OWNERS, "years": a["years"], "games": a["games"],
                        "first": a["first"], "second": a["second"], "third": a["third"],
                        "top2": a["first"] + a["second"], "top3": a["first"] + a["second"] + a["third"],
                        "W": a["W"], "L": a["L"], "T": a["T"],
                        "win_pct": round((a["W"] + 0.5 * a["T"]) / a["games"], 4) if a["games"] else 0,
                        "playoffs": a["playoffs"],
                        "playoff_pct": round(a["playoffs"] / a["years"], 4)})
    alltime.sort(key=lambda x: (not x["active"], -x["win_pct"]))
    act = [x for x in alltime if x["active"]]
    ranks = {x["manager"]: {} for x in act}
    for k in ("years", "games", "first", "second", "third", "top2", "top3", "W", "L",
              "win_pct", "playoffs", "playoff_pct"):
        for x in act:
            better = sum(1 for y in act if (y[k] < x[k] if k == "L" else y[k] > x[k]))
            ranks[x["manager"]][k] = better + 1
    rookies = [o for o in OWNERS if o not in agg]
    years = sorted({s["year"] for s in seasons})
    wins_grid = {y: {} for y in years}
    playoff_grid = {y: {} for y in years}
    for s in seasons:
        wins_grid[s["year"]][s["manager"]] = s["W"]
        playoff_grid[s["year"]][s["manager"]] = int(s["playoffs"])
    games_by_year = {y: max(s["games"] for s in seasons if s["year"] == y) for y in years}
    champs = []
    for y in years:
        pick = lambda k: next((s for s in seasons if s["year"] == y and s[k] == 1), None)
        c, ru, th = pick("first"), pick("second"), pick("third")
        if not c:
            flags.append(f"History {y}: no champion marked")
        champs.append({"year": y, "label": season_label(y),
                       "champion": c["manager"] if c else None,
                       "runner_up": ru["manager"] if ru else None, "third": th["manager"] if th else None,
                       "W": c["W"] if c else None, "L": c["L"] if c else None, "T": c["T"] if c else None,
                       "teams": sum(1 for s in seasons if s["year"] == y)})
    # evolution: seats keep their position year to year
    members = {y: [s["manager"] for s in seasons if s["year"] == y] for y in years}
    members[SEASON] = list(OWNERS)
    all_years = sorted(members)
    last_year = {m: max(y for y in all_years if m in members[y]) for y in all_years for m in members[y]}
    evolution, seats, seen = [], [], set()
    for i, y in enumerate(all_years):
        cur = members[y]
        seats = [s if s in cur else None for s in seats]
        for m in sorted(cur, key=lambda m: (m not in OWNERS, m)):
            if m not in seats:
                if None in seats:
                    seats[seats.index(None)] = m
                else:
                    seats.append(m)
        prev = members[all_years[i - 1]] if i else []
        row = []
        for m in seats:
            if m is None:
                row.append(None)
                continue
            if i == 0:
                t = "Original"
            elif m not in seen:
                t = "New"
            elif m not in prev:
                t = "Returning"
            else:
                t = "Continuing"
            if last_year[m] == y and y != SEASON:
                t = "Departing" if t == "Continuing" else t
            row.append({"manager": m, "type": t})
        seen.update(cur)
        while row and row[-1] is None:
            row.pop()
        evolution.append({"year": y, "label": season_label(y), "seats": row})

    awards = []
    def add(title, key, rows, rev=True, extra=lambda r: "All", min_val=None):
        if not rows:
            return
        top = max(r[key] for r in rows) if rev else min(r[key] for r in rows)
        if min_val is not None and top < min_val:
            return
        who = [r for r in rows if abs(r[key] - top) < 1e-9]
        awards.append({"award": title, "stat": top, "managers": sorted({r["manager"] for r in who}),
                       "years": sorted({str(extra(r)) for r in who})})
    add("Most Titles", "first", alltime)
    add("Most Runner-Up Finishes", "second", alltime)
    add("Most Wins", "W", alltime)
    add("Most Losses", "L", alltime)
    add("Best Win % (5+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 5])
    add("Worst Win % (5+ seasons)", "win_pct", [a for a in alltime if a["years"] >= 5], rev=False)
    add("Most Playoff Appearances", "playoffs", alltime)
    add("Most Top 3 Finishes", "top3", alltime)
    sp = [dict(s, pct=(s["W"] + 0.5 * s["T"]) / s["games"]) for s in seasons if s["games"]]
    add("Best Single Season (Win %)", "pct", sp, extra=lambda r: season_label(r["year"]))
    add("Worst Single Season (Win %)", "pct", sp, rev=False, extra=lambda r: season_label(r["year"]))
    # longest playoff streak
    best = (0, [], [])
    for m in agg:
        run, start = 0, None
        for y in years:
            if playoff_grid[y].get(m):
                run += 1
                start = start or y
                if run > best[0]:
                    best = (run, [m], [f"{season_label(start)} to {season_label(y)}"])
                elif run == best[0]:
                    best[1].append(m); best[2].append(f"{season_label(start)} to {season_label(y)}")
            else:
                run, start = 0, None
    if best[0]:
        awards.append({"award": "Longest Playoff Streak", "stat": best[0],
                       "managers": best[1], "years": best[2]})
    return {"alltime": alltime, "active_ranks": ranks, "rookies": rookies, "years": years,
            "labels": {y: season_label(y) for y in all_years}, "games_by_year": games_by_year,
            "wins_grid": wins_grid, "playoff_grid": playoff_grid, "championships": champs,
            "evolution": evolution, "awards": awards}


# ----------------------------------------------------------------- MAIN
def main():
    flags = []
    swb = load_workbook("STATS_FILE", STATS_URL)
    hwb = load_workbook("HISTORY_FILE", HISTORY_URL)

    weeks = load_weekly(swb["Free For All"], flags)
    pair_opponents(weeks, flags)

    players, ytot = [], {}
    for sheet, scoring, kind in (("SKATER PULL", SKATER_SCORING, "S"), ("GOALIE PULL", GOALIE_SCORING, "G")):
        p, t = read_pull(swb[sheet], scoring, kind, flags)
        for pl in p:
            pl["pts_by_cat"] = score(pl["stats"], scoring)
            pl["points"] = sum(pl["pts_by_cat"].values())
            if not pl["pos"]:
                flags.append(f"Could not read position for '{pl['player']}' ({pl['owner']})")
        players += p
        ytot[kind] = t
    missing = [o for o in OWNERS if o not in {p["owner"] for p in players}]
    if missing:
        flags.append(f"No game log rows found for {', '.join(missing)}")

    official = read_data_pull(swb["Data Pull"]) if "Data Pull" in swb.sheetnames else {}
    recon, status, weeks_entered, verified = reconcile(players, ytot, weeks, OWNERS, official, flags)
    standings, weekly_rows = build_standings(weeks, OWNERS)
    cats, positions = build_analytics(players, OWNERS)
    history = build_history(hwb, flags)

    player_out = sorted(({"owner": p["owner"], "player": p["player"], "team": p["team"],
                          "pos": p["pos"], "elig": p["elig"], "gp": int(p["stats"].get("GP", 0)),
                          "points": r2(p["points"]),
                          "stats": {k: v for k, v in p["stats"].items() if k != "GP"},
                          "pts_by_cat": {k: r2(v) for k, v in p["pts_by_cat"].items()}}
                         for p in players), key=lambda x: -x["points"])

    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    reg_weeks = [w for w in weeks if w <= REG_SEASON_WEEKS]
    out = {"season": SEASON, "season_label": season_label(SEASON), "updated": now,
           "weeks_entered": weeks_entered, "last_reg_week": reg_weeks[-1] if reg_weeks else 0,
           "verified_through": verified,
           "reg_season_weeks": REG_SEASON_WEEKS, "playoff_weeks": PLAYOFF_WEEKS,
           "playoff_teams": PLAYOFF_TEAMS, "owners": OWNERS, "recon_status": status,
           "recon": recon, "flags": flags, "standings": standings, "weekly": weekly_rows,
           "categories": cats, "category_groups": CATEGORY_GROUPS, "cat_labels": CAT_LABELS,
           "positions": positions, "players": player_out, "history": history,
           "scoring": {"skaters": SKATER_SCORING, "goalies": GOALIE_SCORING},
           "yahoo_url": YAHOO_LEAGUE_URL}
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(os.path.join(OUT_DIR, "td_data.json"), "w") as f:
        json.dump(out, f, indent=1, default=float)

    lines = [f"# Type Disgusting Reconciliation: {status}",
             f"Updated {now} | Weeks entered: {weeks_entered} | Yahoo totals verified through week {verified}", "",
             "| Team | Free For All total | Yahoo official | Game log (daily) | Current week so far |",
             "|---|---|---|---|---|"]
    for r in recon:
        off = f"{r['yahoo_official']:.2f}" if r["yahoo_official"] is not None else "n/a"
        lines.append(f"| {r['team']} | {r['weekly_total']:.2f} | {off} | {r['game_log']:.2f} | {r['in_progress']:.2f} |")
    lines += ["", "## Flags"] + ([f"- {f}" for f in flags] or ["- None"])
    with open(os.path.join(OUT_DIR, "recon_report.md"), "w") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines))
    # never fail the run on data flags; the site shows the status badge instead


if __name__ == "__main__":
    main()
